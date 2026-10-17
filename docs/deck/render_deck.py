"""Render deck.html -> slides/NN.png, ../cover.png (slide 1), Foreman_deck.pdf and Foreman_deck.pptx (with speaker notes)."""
from pathlib import Path

from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).parent
NOTES = [
    "Foreman finds a defect, explains why it happened, and drafts the repair. We're in the Intelligent Industry track, and everything runs on a single AMD Ryzen laptop processor, CPU only.",
    "Two costs: scrap until someone notices, then downtime until someone understands. Vision systems stop at 'bad part'. The link from symptom to cause to fix lives in a manual and in experienced people, and the camera feed often can't go to a cloud API.",
    "Real output from the live app, not a mock-up. The lubricant fault is running: the camera flags a discoloured nut, the heatmap boxes the stain, and Gemma 4 names a lubrication system fault. Foreman's own consistency check shows camera and sensors agree via manual section 5.3 to 4.3. Gemma 4 E2B ran CPU-only on the Ryzen for this.",
    "Left to right. Two steps are neural: Detect with PatchCore and Reason with Gemma 4, both on the AMD Ryzen CPU. Sensors use plain statistics, retrieval is checkable, every citation is verified, and a human always approves.",
    "Real faults show up in the machine before they show up in the parts. Foreman acts on the sensor drift first, then upgrades the same incident when the camera confirms it, so the floor gets one incident and one work order, not a flood of alerts.",
    "Everything runs on the laptop's AMD Ryzen 7: the detector and Gemma 4, CPU only, with the NVIDIA card switched off. That's the shape of a real edge PC next to a press, and camera images never leave the plant. A ROCm and vLLM build for the Instinct MI300X is in the repo for plant-wide scale; it's built but we haven't run it yet.",
    "Every number here is measured. The detector threshold is set from good parts only. The live-model eval is small, three images per fault, and we say so. Camera images are real; sensors are simulated and the manual is fictional.",
    "What it takes to run this for real: the plant's own manuals, the real sensor historian, and approved work orders sent into the CMMS. These are roadmap items, not built yet. Code is on GitHub.",
]

with sync_playwright() as p:
    b = p.chromium.launch(channel="chrome", headless=True)
    pg = b.new_page(viewport={"width": 1920, "height": 1080}, device_scale_factor=1)
    pg.goto((HERE / "deck.html").as_uri())
    pg.evaluate("() => document.fonts.ready")
    pg.wait_for_timeout(800)
    out = HERE / "slides"
    out.mkdir(exist_ok=True)
    slides = pg.locator("section.slide")
    n = slides.count()
    for i in range(n):
        slides.nth(i).screenshot(path=str(out / f"{i + 1:02d}.png"))
    pg.pdf(path=str(HERE / "Foreman_deck.pdf"), width="1920px", height="1080px", print_background=True,
           margin={"top": "0", "right": "0", "bottom": "0", "left": "0"})
    b.close()
(HERE.parent / "cover.png").write_bytes((out / "01.png").read_bytes())

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
for i in range(n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(out / f"{i + 1:02d}.png"), 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = NOTES[i]
prs.save(str(HERE / "Foreman_deck.pptx"))
print(n, "slides")
